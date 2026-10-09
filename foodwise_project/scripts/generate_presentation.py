from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from foodwise.calculations import household_result
from foodwise.data import (
    CANONICAL_FOODS,
    DEFAULT_PRICE_FILE,
    filter_prices,
    latest_market_prices,
    load_price_data,
    market_coverage,
)

ASSETS = ROOT / "assets"
OUTPUT = ROOT / "presentation" / "FOODWISE_Project_Presentation.pptx"
NAVY = "#143B3A"
GREEN = "#19745F"
MINT = "#DDEFE7"
GOLD = "#E2B447"
PALE = "#F4F7F5"
WHITE = "#FFFFFF"
INK = "#243A35"
MUTED = "#62756F"
RED = "#A04D3F"


def rgb(hex_value: str) -> RGBColor:
    value = hex_value.lstrip("#")
    return RGBColor(int(value[:2], 16), int(value[2:4], 16), int(value[4:], 16))


def put_text(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    *,
    size: int = 18,
    color: str = INK,
    bold: bool = False,
    align=None,
    font: str = "Aptos",
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.04)
    frame.margin_bottom = Inches(0.02)
    para = frame.paragraphs[0]
    para.text = text
    para.font.name = font
    para.font.size = Pt(size)
    para.font.bold = bold
    para.font.color.rgb = rgb(color)
    if align is not None:
        para.alignment = align
    return shape


def add_background(prs: Presentation, slide_no: int, title: str, kicker: str = ""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(PALE)
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = rgb(GREEN)
    accent.line.fill.background()
    if kicker:
        put_text(slide, 0.62, 0.30, 11.95, 0.26, kicker.upper(), size=9, color=GREEN, bold=True)
    put_text(slide, 0.62, 0.65, 12.0, 0.65, title, size=27, color=NAVY, bold=True, font="Aptos Display")
    put_text(slide, 0.65, 7.14, 10.8, 0.2, "FOODWISE  |  Public-data concept · Kenya · October 2026", size=8, color=MUTED)
    put_text(slide, 12.15, 7.08, 0.48, 0.24, f"{slide_no:02d}", size=9, color=GREEN, bold=True, align=PP_ALIGN.RIGHT)
    return slide


def add_card(slide, x, y, w, h, title, body, *, fill=WHITE, accent=GREEN, title_size=15, body_size=12):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb("#E0E9E4")
    put_text(slide, x + 0.2, y + 0.18, w - 0.38, 0.38, title, size=title_size, color=accent, bold=True)
    put_text(slide, x + 0.2, y + 0.68, w - 0.4, h - 0.8, body, size=body_size, color=INK)
    return shape


def add_metric(slide, x, y, w, value, label, note):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(2.15)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(WHITE)
    shape.line.color.rgb = rgb("#DFE9E3")
    put_text(slide, x + 0.13, y + 0.25, w - 0.26, 0.56, value, size=26, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    put_text(slide, x + 0.17, y + 0.94, w - 0.34, 0.66, label, size=12, color=INK, bold=True, align=PP_ALIGN.CENTER)
    put_text(slide, x + 0.2, y + 1.67, w - 0.4, 0.35, note, size=8, color=MUTED, align=PP_ALIGN.CENTER)


def prepare_charts(prices: pd.DataFrame):
    ASSETS.mkdir(exist_ok=True)
    recent_start = date(2024, 10, 1)
    latest_date = prices["date"].max().date()
    recent = filter_prices(
        prices,
        start=recent_start,
        end=latest_date,
        price_type="Retail",
    )
    target_recent = recent[
        recent["commodity"].isin(CANONICAL_FOODS + ("Sugar",))
    ]
    coverage = market_coverage(
        target_recent,
        start=recent_start,
        end=latest_date,
        price_type=None,
    ).head(9)
    fig, ax = plt.subplots(figsize=(8.1, 3.4), dpi=180)
    plotted = coverage.iloc[::-1]
    bars = ax.barh(plotted["market"], plotted["commodities"], color="#19745F")
    ax.set_xlim(0, max(9, int(plotted["commodities"].max()) + 1))
    ax.set_xlabel("Distinct target foods observed since Oct 2024")
    ax.set_title("Recent coverage by market varies", loc="left", color="#143B3A", fontsize=13, pad=12)
    ax.grid(axis="x", alpha=0.2)
    ax.spines[["top", "right", "left"]].set_visible(False)
    ax.tick_params(axis="y", length=0, labelsize=8)
    ax.bar_label(bars, padding=4, fontsize=8)
    fig.tight_layout()
    coverage_path = ASSETS / "recent_market_food_coverage.png"
    fig.savefig(coverage_path, transparent=False, facecolor="white")
    plt.close(fig)

    trend_market = "IFO (Daadab)"
    trend_items = ["Maize flour", "Beans (dry)", "Milk (cow, fresh)"]
    selected = filter_prices(
        prices,
        start=recent_start,
        end=latest_date,
        markets=[trend_market],
        commodities=trend_items,
        price_type="Retail",
    )
    fig, ax = plt.subplots(figsize=(8.1, 3.4), dpi=180)
    colors_by_item = {"Maize flour": "#19745F", "Beans (dry)": "#D19A22", "Milk (cow, fresh)": "#3D6EA8"}
    for item in trend_items:
        item_rows = selected[selected["commodity"] == item]
        if item_rows.empty:
            continue
        monthly = (
            item_rows.groupby(["month", "analysis_unit"], as_index=False)
            .agg(median=("price_per_analysis_unit", "median"))
        )
        for unit in monthly["analysis_unit"].unique():
            unit_rows = monthly[monthly["analysis_unit"] == unit]
            ax.plot(
                unit_rows["month"],
                unit_rows["median"],
                marker="o",
                linewidth=2,
                markersize=3,
                color=colors_by_item[item],
                label=f"{item} · {unit}",
            )
    ax.set_title("Observed monthly medians · IFO (Daadab)", loc="left", color="#143B3A", fontsize=13, pad=12)
    ax.set_ylabel("KES per normalized unit")
    ax.grid(axis="y", alpha=0.2)
    ax.spines[["top", "right"]].set_visible(False)
    ax.legend(frameon=False, ncol=2, fontsize=8, loc="upper left")
    fig.autofmt_xdate(rotation=30)
    fig.tight_layout()
    trend_path = ASSETS / "ifo_price_trends.png"
    fig.savefig(trend_path, transparent=False, facecolor="white")
    plt.close(fig)
    return coverage, coverage_path, trend_path, trend_market, selected


def build_deck():
    prices = load_price_data(DEFAULT_PRICE_FILE)
    coverage, coverage_chart, trend_chart, scenario_market, scenario_source = prepare_charts(prices)
    latest = latest_market_prices(
        prices,
        market=scenario_market,
        commodities=list(CANONICAL_FOODS),
        price_type="Retail",
    )
    quantity = {
        "Maize flour": 1.5,
        "Beans (dry)": 0.35,
        "Kale": 0.75,
        "Milk (cow, fresh)": 1.75,
        "Oil (vegetable)": 0.25,
        "Potatoes (Irish)": 0.75,
        "Rice": 0.35,
    }
    liquid = {"Milk (cow, fresh)", "Oil (vegetable)"}
    latest_prices = {
        (row["commodity"], row["analysis_unit"]): row
        for _, row in latest.iterrows()
    }
    basket_weekly = 0.0
    priced = 0
    stale_items: list[str] = []
    scenario_cutoff = prices["date"].max()
    for item, amount in quantity.items():
        expected_unit = "KES/L" if item in liquid else "KES/kg"
        row = latest_prices.get((item, expected_unit))
        if row is not None:
            basket_weekly += amount * float(row["price_per_analysis_unit"])
            priced += 1
            if (scenario_cutoff - row["date"]).days > 90:
                stale_items.append(item)
    scenario = household_result(
        weekly_basket_per_person=basket_weekly,
        household_size=4,
        monthly_income=25_000,
        priced_items=priced,
        requested_items=len(quantity),
    )

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "FOODWISE | Food affordability intelligence for Kenya"
    prs.core_properties.subject = "A public-data prototype and county-pilot proposal"
    prs.core_properties.author = "FOODWISE"

    # 1 — Cover
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(NAVY)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.68), Inches(0.92), Inches(0.12), Inches(4.95))
    band.fill.solid()
    band.fill.fore_color.rgb = rgb(GOLD)
    band.line.fill.background()
    put_text(slide, 1.08, 1.22, 10.9, 0.4, "KENYA · PUBLIC-DATA CONCEPT", size=11, color="#B8D9CD", bold=True)
    put_text(slide, 1.05, 1.86, 11, 1.12, "FOODWISE", size=52, color=WHITE, bold=True, font="Aptos Display")
    put_text(slide, 1.1, 3.02, 10.7, 0.9, "Making healthy-food affordability visible", size=27, color="#CDE8DE", bold=True)
    put_text(slide, 1.1, 4.3, 10.7, 0.78, "Real market-price signals. Transparent household scenarios. Honest data limits.", size=18, color=WHITE)
    put_text(slide, 1.1, 6.18, 10.9, 0.42, "[Presenter name]  ·  [Email / phone]", size=13, color="#B8D9CD")
    put_text(slide, 12.2, 7.1, 0.45, 0.2, "01", size=9, color="#B8D9CD", align=PP_ALIGN.RIGHT)

    # 2 — Why it matters
    slide = add_background(prs, 2, "The affordability challenge is real—and measurable.", "01 / THE CASE")
    add_metric(slide, 0.72, 1.72, 3.8, "43.5–43.9M", "People unable to afford a healthy diet", "FAO briefing figures · exact series/year to confirm")
    add_metric(slide, 4.77, 1.72, 3.8, "KSh 189–582", "Estimated daily healthy-diet cost", "FAO / CoAHD briefing range · method-dependent")
    add_metric(slide, 8.82, 1.72, 3.8, "9.5%", "Food & Non-Alcoholic Beverages inflation", "KNBS · year-on-year · September 2026")
    put_text(slide, 0.9, 4.48, 11.6, 0.52, "The build starts from public data—not from pretending these sources are interchangeable.", size=18, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    put_text(slide, 1.25, 5.34, 10.8, 0.72, "KNBS CPI is a national index. WFP provides market observations. FAO CoAHD uses its own diet-cost method.", size=15, color=MUTED, align=PP_ALIGN.CENTER)
    put_text(slide, 1.1, 6.32, 11.0, 0.42, "FAO figures above are included from the project brief; confirm the exact dataset, year and denominator before public claims.", size=9, color=MUTED, align=PP_ALIGN.CENTER)

    # 3 — Source and freshness
    slide = add_background(prs, 3, "The prototype is built on a real WFP Kenya price file.", "02 / EVIDENCE BASE")
    metrics = [
        ("27,897", "WFP source rows"),
        ("2006–2026", "Recorded date range"),
        ("1,689", "Recent target-food rows since Oct 2024"),
        ("15 Sep 2026", "Latest WFP source date"),
    ]
    for i, (value, label) in enumerate(metrics):
        x = 0.78 + i * 3.13
        put_text(slide, x, 1.5, 2.8, 0.48, value, size=23, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
        put_text(slide, x, 2.03, 2.8, 0.38, label, size=11, color=MUTED, align=PP_ALIGN.CENTER)
    slide.shapes.add_picture(str(coverage_chart), Inches(0.82), Inches(2.65), width=Inches(7.1))
    add_card(slide, 8.1, 2.75, 4.4, 2.1, "What the chart means", "Coverage counts distinct target foods in recent retail observations. It is a snapshot of data availability—not evidence that one county is more expensive or more food insecure.", title_size=15, body_size=12)
    put_text(slide, 8.25, 5.17, 4.1, 0.78, "Current WFP coverage includes markets in Garissa and Turkana, including refugee-camp markets.", size=12, color=INK)
    put_text(slide, 8.25, 6.0, 4.1, 0.55, "Not a Nairobi/Nyeri or county-representative retail panel.", size=12, color=RED, bold=True)

    # 4 — Live price evidence
    slide = add_background(prs, 4, "Price signals move by food, month and market.", "03 / PRICE INTELLIGENCE")
    slide.shapes.add_picture(str(trend_chart), Inches(0.72), Inches(1.55), width=Inches(7.95))
    add_card(slide, 8.95, 1.7, 3.7, 1.28, "Monthly median", "Uses reported WFP Retail observations; explicit package units normalized per kg / litre.", title_size=14, body_size=10)
    add_card(slide, 8.95, 3.13, 3.7, 1.28, "One market example", f"{scenario_market} · a source market, not a proxy for all households.", title_size=14, body_size=10)
    add_card(slide, 8.95, 4.56, 3.7, 1.28, "Interpret with care", "Markets and record counts vary over time; this is not KNBS CPI or a causal diet outcome.", fill="#FFF6E2", accent="#9A6D16", title_size=14, body_size=10)
    put_text(slide, 0.88, 6.42, 11.2, 0.35, "Series shown: maize flour, dry beans and fresh cow milk. Exact dates, units and flags are available in the app.", size=10, color=MUTED)

    # 5 — Product
    slide = add_background(prs, 5, "Three linked modules turn prices into questions people can act on.", "04 / THE PRODUCT")
    add_card(slide, 0.75, 1.72, 3.85, 2.53, "01  Price Intelligence", "Browse food prices by market and month. Compare like-for-like units. Export the filtered source observations.", title_size=17, body_size=13)
    add_card(slide, 4.75, 1.72, 3.85, 2.53, "02  Household Pressure Scenario", "Change the basket quantities, household size and income. See the priced-item cost share and the observations behind it.", title_size=17, body_size=13)
    add_card(slide, 8.75, 1.72, 3.85, 2.53, "03  Diet Gap Readiness", "Use expert-reviewed nutrient composition and population-specific references—without fabricating nutrient data from prices.", title_size=17, body_size=13)
    put_text(slide, 1.0, 5.12, 11.2, 0.68, "The key design choice: missing values remain visible as missing.", size=22, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    put_text(slide, 1.0, 5.9, 11.2, 0.45, "No silent zero prices · no invented nutrient scores · no CPI/market-price conflation", size=13, color=MUTED, align=PP_ALIGN.CENTER)

    # 6 — Household scenario
    slide = add_background(prs, 6, "A household scenario is now interactive—not a claimed healthy-diet cost.", "05 / HOUSEHOLD SIMULATOR")
    weekly_text = f"KSh {scenario.weekly_per_person:,.0f}"
    monthly_text = f"KSh {scenario.monthly_household:,.0f}"
    share_text = f"{scenario.income_share:.1%}" if scenario.income_share is not None else "Income required"
    freshness_note = (
        f"{len(stale_items)} observation(s) older than 90 days"
        if stale_items
        else "All matching observations within 90 days"
    )
    add_metric(slide, 0.8, 1.66, 3.7, weekly_text, "Weekly scenario per person", f"{priced} of {len(quantity)} items priced · {freshness_note}")
    add_metric(slide, 4.8, 1.66, 3.7, monthly_text, "Monthly household scenario", "4 people · illustrative quantities")
    add_metric(slide, 8.8, 1.66, 3.7, share_text, "Share of KSh 25,000 monthly income", "Prototype calculation · not a validated risk band")
    add_card(slide, 0.9, 4.3, 5.5, 1.55, "Editable in the app", "Food quantities · market · household size · income · watch / critical thresholds.", title_size=15, body_size=12)
    add_card(slide, 6.75, 4.3, 5.5, 1.55, "Why the caveat matters", "The basket is a demonstration scenario. It is not an official FAO CoAHD basket or nutrition recommendation.", fill="#FFF6E2", accent="#9A6D16", title_size=15, body_size=12)
    stale_copy = (
        "One or more source prices are more than 90 days older than the selected cutoff; check each date in the app."
        if stale_items
        else "Prices use latest matching WFP Retail records up to 15 Sep 2026."
    )
    put_text(slide, 1.0, 6.28, 11.2, 0.34, stale_copy, size=10, color=MUTED, align=PP_ALIGN.CENTER)

    # 7 — Diet gap
    slide = add_background(prs, 7, "Diet-gap analysis requires nutrition data—not just price.", "06 / RESPONSIBLE SCOPE")
    stages = [
        ("1", "Choose foods & quantities", "Same scenario inputs used by the household calculator."),
        ("2", "Provide sourced composition", "Nutrient amounts per kg / litre / source unit, matched to foods."),
        ("3", "Set a population reference", "Expert-selected age / life-stage reference and source year."),
        ("4", "Compare transparently", "Estimate contribution and remaining amount; show sources."),
    ]
    for index, (number, heading, body) in enumerate(stages):
        y = 1.48 + index * 1.15
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.95), Inches(y), Inches(0.48), Inches(0.48))
        circle.fill.solid()
        circle.fill.fore_color.rgb = rgb(GREEN)
        circle.line.fill.background()
        put_text(slide, 0.95, y + 0.1, 0.48, 0.22, number, size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        put_text(slide, 1.7, y, 3.2, 0.38, heading, size=16, color=NAVY, bold=True)
        put_text(slide, 5.0, y + 0.02, 7.2, 0.57, body, size=13, color=INK)
    put_text(slide, 1.05, 6.28, 11.4, 0.42, "No composition dataset is bundled today; the interface waits for reviewed inputs rather than inventing outputs.", size=12, color=RED, bold=True, align=PP_ALIGN.CENTER)

    # 8 — How it works
    slide = add_background(prs, 8, "From official public files to an auditable live prototype.", "07 / HOW IT WORKS")
    blocks = [
        ("INGEST", "WFP price CSV\nWFP market coordinates\nKNBS CPI reference"),
        ("VALIDATE", "Required fields\nKES & positive values\nComparable units"),
        ("ANALYSE", "Monthly medians\nLatest market price\nEditable scenarios"),
        ("SHARE", "Interactive Streamlit app\nFiltered CSV export\nPowerPoint showcase"),
    ]
    for index, (heading, body) in enumerate(blocks):
        x = 0.62 + index * 3.2
        add_card(slide, x, 2.05, 2.75, 2.35, heading, body, title_size=14, body_size=13)
        if index < len(blocks) - 1:
            put_text(slide, x + 2.77, 2.95, 0.42, 0.42, "›", size=28, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    put_text(slide, 0.95, 5.15, 11.5, 0.56, "Auditable methods live in the project: source CSVs · unit conversions · tests · limitations · reproducible setup.", size=15, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    put_text(slide, 1.0, 5.95, 11.2, 0.4, "Run locally with Python + Streamlit. The app bundles the data and can filter an updated WFP-style CSV.", size=12, color=MUTED, align=PP_ALIGN.CENTER)

    # 9 — Pilot roadmap
    slide = add_background(prs, 9, "The next step is a validated county pilot—not a bigger claim.", "08 / NEXT 30–60 DAYS")
    roadmap = [
        ("01 · Validate", "County / market coverage\nNutrition expert reviews basket\nConfirm CoAHD benchmarks"),
        ("02 · Co-design", "Agree household groups\nAgree reference period\nDefine usable outputs"),
        ("03 · Pilot", "Collect matched local prices\nTest with households & partners\nMeasure coverage and usability"),
        ("04 · Decide", "Document limits and learnings\nPrioritize additional markets\nScale only if evidence supports it"),
    ]
    for index, (heading, body) in enumerate(roadmap):
        x = 0.72 + index * 3.15
        add_card(slide, x, 1.78, 2.78, 2.7, heading, body, fill=WHITE, title_size=15, body_size=12)
    put_text(slide, 1.05, 5.28, 11.1, 0.65, "Pilot partners needed: one county team · a nutrition reviewer · trusted market-price enumerators.", size=16, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    put_text(slide, 1.2, 6.15, 10.8, 0.42, "Success starts with credible local coverage and a validated basket, not dashboard complexity.", size=12, color=MUTED, align=PP_ALIGN.CENTER)

    # 10 — Close
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(NAVY)
    put_text(slide, 1.0, 0.95, 11.35, 0.42, "HEALTHY FOOD SHOULD NOT BE A LUXURY.", size=12, color="#B8D9CD", bold=True, align=PP_ALIGN.CENTER)
    put_text(slide, 1.05, 1.9, 11.2, 1.45, "FOODWISE makes the barriers visible so we can act on them.", size=31, color=WHITE, bold=True, align=PP_ALIGN.CENTER, font="Aptos Display")
    put_text(slide, 2.0, 4.15, 9.3, 0.62, "Seeking county partners · nutrition mentors · a pilot opportunity", size=18, color="#CDE8DE", align=PP_ALIGN.CENTER)
    put_text(slide, 2.0, 5.65, 9.3, 0.42, "[Presenter name]  ·  [Email]  ·  [Phone]", size=14, color=WHITE, align=PP_ALIGN.CENTER)
    put_text(slide, 12.2, 7.1, 0.45, 0.2, "10", size=9, color="#B8D9CD", align=PP_ALIGN.RIGHT)

    prs.save(OUTPUT)
    print(f"Presentation: {OUTPUT}")
    print(f"Scenario: {scenario_market}; {priced}/{len(quantity)} foods with prices")
    print(f"Weekly per person: KSh {scenario.weekly_per_person:,.2f}")
    print(f"Monthly household: KSh {scenario.monthly_household:,.2f}")
    print(f"Income share: {scenario.income_share:.1%}" if scenario.income_share is not None else "Income share unavailable")
    print(f"Recent coverage markets: {len(coverage)} shown")


if __name__ == "__main__":
    build_deck()
